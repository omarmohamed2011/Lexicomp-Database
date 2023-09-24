import re

from docx import Document

import logging


def get_drug_interactions(paragraphs):
    found = False
    text = ""
    for paragraph in paragraphs:
        if paragraph.style.name.startswith('Heading 1') and paragraph.text == 'Drug Interactions':
            found = True
            continue
        if found and paragraph.style.name != "Heading 1":
            text += paragraph.text
            text += '\n\n'
        elif found and paragraph.style.name == "Heading 1":
            break
    return text


def get_paragraphs_lst(paragraphs):
    paragraph_list = []
    for paragraph in paragraphs:
        paragraph_list.append(paragraph.text)
    return paragraph_list


def get_drug_interactions2(paragraphs):
    found = False
    text = []
    for paragraph in paragraphs:
        if paragraph.style.name.startswith('Heading 1') and paragraph.text == 'Drug Interactions':
            found = True
            continue
        if found and paragraph.style.name != "Heading 1":
            text.append(paragraph.text)
        elif found and paragraph.style.name == "Heading 1":
            break
    return text


def get_heading_text(paragraphs, heading_name):
    found = False
    text = ""
    for paragraph in paragraphs:
        if paragraph.style.name.startswith('Heading 1') and (heading_name in paragraph.text):
            found = True
            continue
        if found and paragraph.style.name != "Heading 1":
            text += paragraph.text
            text += '\n\n'
        elif found and paragraph.style.name == "Heading 1":
            break
    return text


def get_title(document):
    # "3/16/23, 7:18 PM"
    time_regex = "\d\d?/\d\d?/23, \d\d?:\d\d? ((PM)|(AM))"
    regex_list = ["\(Lab Tests and Diagnostic Procedures\)", "\(Lexi.*?\)"]
    regex = "|".join([f"({reg})" for reg in regex_list])

    if re.search(time_regex, document.paragraphs[0].text) is not None:
        title = document.paragraphs[2].text
    else:
        title = document.paragraphs[0].text

    for paragraph in document.paragraphs:
        txt = paragraph.text
        result = re.search(regex, txt)
        if result is not None:
            title = txt
            break

    return title


def read_doc(path):
    document = Document(path)
    return document


def clear_logger():
    for handler in logging.root.handlers[:]:
        logging.root.removeHandler(handler)


def words_count(document):
    cnt = 0
    for paragraph in document.paragraphs:
        txt = paragraph.text
        cnt += len(txt.split())
    return cnt


